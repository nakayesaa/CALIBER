#!/usr/bin/env python3
"""Verify HE sources and build a weekly-native evidence investigation artifact."""

from __future__ import annotations

import argparse
import sys
from datetime import datetime
from pathlib import Path
from zoneinfo import ZoneInfo

import yaml
from openpyxl import load_workbook
from pptx import Presentation

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from services.api.app.schemas.alerts import AlertEvent, AlertStateTransition
from services.api.app.schemas.api import AssetSummary, InvestigationEvidenceEvent
from services.api.app.schemas.equipment import (
    EquipmentAssessment,
    EquipmentInvestigation,
    EquipmentOperatingState,
    EquipmentPoint,
    EquipmentSignal,
    EquipmentSourceAction,
)
from services.api.app.services.file_io import atomic_write_json, file_sha256

ZONE = ZoneInfo("Asia/Jakarta")
ALERT_ID = "alert-asset-he-3301-0001"


def timestamp(value: str, pattern: str = "%Y-%m-%d") -> datetime:
    return datetime.strptime(value, pattern).replace(tzinfo=ZONE)


def assess(signals: list[EquipmentSignal], policy: dict) -> list[EquipmentAssessment]:
    """Assess inclusive limits, persistent breaches and corroborated weekly trend."""
    results = []
    alarm_streak = trend_streak = 0
    baseline = {s.key: s.points[0].value for s in signals}
    for index in range(len(signals[0].points)):
        alarm, trip, progress, worsening = [], [], [], 0
        for signal in signals:
            value = signal.points[index].value
            sign = 1 if signal.direction == "HIGH" else -1
            if sign * (value - signal.alarm_limit) >= 0:
                alarm.append(signal.key)
            if sign * (value - signal.trip_limit) >= 0:
                trip.append(signal.key)
            distance = sign * (signal.trip_limit - baseline[signal.key])
            progress.append(max(0, min(1, sign * (value - baseline[signal.key]) / distance)))
            if index and sign * (value - signal.points[index - 1].value) > 0:
                worsening += 1
        alarm_streak = alarm_streak + 1 if alarm else 0
        trend_streak = trend_streak + 1 if worsening >= policy["watch_worsening_signals"] else 0
        if trip:
            state, reason = "CRITICAL", "At least one engineering trip limit reached"
        elif (
            len(alarm) >= policy["high_minimum_alarm_signals"]
            and alarm_streak >= policy["warning_consecutive_weeks"]
        ):
            state, reason = "HIGH", "Persistent alarm with at least three corroborating signals"
        elif alarm_streak >= policy["warning_consecutive_weeks"]:
            state, reason = "WARNING", "Alarm breach persists for two weekly observations"
        elif alarm or (
            trend_streak >= policy["watch_consecutive_weeks"]
            and max(progress) * 100 >= policy["watch_minimum_progress"]
        ):
            state, reason = (
                "WATCH",
                "Alarm boundary or repeated multi-signal worsening requires review",
            )
        else:
            state, reason = "NORMAL", "No sustained adverse trend or alarm breach"
        point = signals[0].points[index]
        results.append(
            EquipmentAssessment(
                timestamp=point.timestamp,
                state=state,
                score=round(max(progress) * 100, 2),
                breached_signals=alarm,
                reason=reason,
                source_status=point.source_status,
            )
        )
    return results


def slide_text(presentation, number: int) -> list[str]:
    return [
        shape.text.strip()
        for shape in presentation.slides[number - 1].shapes
        if shape.has_text_frame and shape.text.strip()
    ]


def source_actions(presentation) -> list[EquipmentSourceAction]:
    result = []
    for number, headings in [
        (9, ["Corrective Action", "Pro-Active Action"]),
        (10, ["Preventive Action"]),
    ]:
        texts = slide_text(presentation, number)
        for heading in headings:
            start = texts.index(heading) + (4 if number == 9 else 3)
            for offset in range(2):
                cells = texts[start + offset * 5 : start + (offset + 1) * 5]
                code = cells[0]
                title, due, owner = cells[1:4] if number == 9 else cells[2:5]
                result.append(
                    EquipmentSourceAction(
                        id=f"he-source-action-{len(result) + 1}",
                        title=title,
                        action_type="CORRECTIVE"
                        if heading == "Corrective Action"
                        else "PREVENTIVE",
                        owner=owner,
                        due_date=datetime.strptime(due, "%d-%b-%Y").date(),
                        status=cells[4] if number == 9 else "Not reported",
                        source_reference=f"he_rca:slide:{number}:{heading}:row:{offset + 1}:{code}",
                    )
                )
    return result


def ingest(root: Path) -> EquipmentInvestigation:
    manifest = yaml.safe_load((root / "data/catalog/he_3301_sources.yaml").read_text())
    policy = yaml.safe_load((root / "data/catalog/he_3301_detection.yaml").read_text())
    paths = {}
    for source in manifest["sources"]:
        path = root / source["local_path"]
        if path.stat().st_size != source["size_bytes"] or file_sha256(path) != source["sha256"]:
            raise ValueError(f"Immutable source checksum mismatch: {source['source_key']}")
        paths[source["source_key"]] = path
    performance = load_workbook(paths["he_performance"], data_only=True)
    info = {row[0]: row[1] for row in performance["Equipment Info"].values if row[0]}
    if info["Equipment Tag"] != "HE-3301":
        raise ValueError("Unexpected equipment identity")
    weekly = list(performance["Condition History"].values)[1:]
    if len(weekly) != 26:
        raise ValueError("Expected 26 weekly source observations")
    signals = []
    for spec in policy["signals"]:
        signals.append(
            EquipmentSignal(
                **{key: value for key, value in spec.items() if key != "column"},
                cadence="WEEKLY",
                source_key="he_performance",
                points=[
                    EquipmentPoint(
                        timestamp=timestamp(row[1]),
                        value=row[spec["column"]],
                        source_status=row[6],
                        source_reference=f"he_performance:Condition History:row:{i}",
                    )
                    for i, row in enumerate(weekly, 2)
                ],
            )
        )
    assessments = assess(signals, policy)
    production = load_workbook(paths["he_production"], data_only=True)
    hourly = list(production["Sheet2"].values)
    expected = (
        "Timestamp",
        "HE3301_FEED",
        "HE3301_DISP",
        "HE3301_VIB",
        "HE3301_TEMP",
        "HE3301_AMP",
        "PLANT_RATE",
        "RUN_STATUS",
    )
    if tuple(hourly[0]) != expected or len(hourly) != 721:
        raise ValueError("Unexpected hourly source schema or row count")
    for key, label, column in [
        ("feed_rate", "Equipment feed rate", 1),
        ("plant_rate", "Source-reported ZCU production rate", 6),
    ]:
        signals.append(
            EquipmentSignal(
                key=key,
                label=label,
                unit="t/h",
                cadence="HOURLY",
                source_key="he_production",
                points=[
                    EquipmentPoint(
                        timestamp=timestamp(row[0], "%Y-%m-%d %H:%M:%S"),
                        value=row[column],
                        source_reference=f"he_production:Sheet2:row:{i}",
                    )
                    for i, row in enumerate(hourly[1:], 2)
                ],
            )
        )
    states = [
        EquipmentOperatingState(
            timestamp=timestamp(row[0], "%Y-%m-%d %H:%M:%S"),
            state=row[7],
            source_reference=f"he_production:Sheet2:row:{i}",
        )
        for i, row in enumerate(hourly[1:], 2)
    ]
    transitions, previous = [], "NORMAL"
    for assessment in assessments:
        if assessment.state != previous:
            transitions.append(
                AlertStateTransition(
                    alert_id=ALERT_ID,
                    timestamp=assessment.timestamp.isoformat(),
                    previous_state=previous,
                    new_state=assessment.state,
                    reason=assessment.reason,
                )
            )
            previous = assessment.state
    first = next(a for a in assessments if a.state != "NORMAL")
    opened = next(a for a in assessments if a.state in {"WARNING", "HIGH", "CRITICAL"})
    peak = max(assessments, key=lambda a: a.score)
    recovered = next(a for a in assessments if a.timestamp > peak.timestamp and a.state == "NORMAL")
    presentation = Presentation(paths["he_rca"])
    texts = slide_text(presentation, 3)
    events = []
    for index, (title, kind) in enumerate(
        [
            ("Reported pre-isolation deterioration", "SENSOR"),
            ("Rate reduction and isolation", "ACTION"),
            ("Bundle deposit inspection", "INSPECTION"),
            ("Hydro-jet cleaning and return to service", "ACTION"),
        ]
    ):
        date_text, detail = texts[3 + index * 2 : 5 + index * 2]
        occurred = timestamp(date_text, "%d-%b-%Y %H:%M" if " " in date_text else "%d-%b-%Y")
        events.append(
            InvestigationEvidenceEvent(
                event_id=f"he-evidence-{index + 1}",
                occurred_at=occurred,
                kind=kind,
                title=title,
                detail=detail,
                source_reference=f"he_rca:slide:3:event:{index + 1}",
                source_grade="RCA_REPORTED_DATE_ONLY" if index == 0 else "RCA_REPORTED",
            )
        )
    summary = {row[0]: row[1] for row in performance["Performance Summary"].values if row[0]}
    issues = [
        "Asia/Jakarta timezone assumed; weekly date-only measurements do not establish an hourly event time.",
        "HE3301_DISP, HE3301_VIB, HE3301_TEMP and HE3301_AMP are unmapped for this static exchanger and excluded from assessment.",
        "RUN_STATUS is a source operating-state record; tag mapping is unconfirmed for a static exchanger.",
        "KO and HE PLANT_RATE scopes are unconfirmed; HE rate is kept source-specific and never summed with KO.",
        "RCA inspection, upstream filter, tube velocity and shell inlet findings lack original attached measurement/inspection records.",
        "Equipment Info design-life and vibration-route text appears copied from rotating equipment; excluded from HE decision logic.",
        "12 March ALARM is preserved as source status; derived WATCH uses inclusive boundary and two-week warning persistence.",
        "RCA 21 May 21:00 return-to-service conflicts with hourly OFF at 21:00 and ON at 22:00.",
        f"Source-reported downtime is {summary['Total Downtime (hours)']:g} h; hourly source contains {sum(s.state == 'OFF' for s in states)} OFF samples. Sample count does not replace elapsed RCA duration.",
        "Engineering score measures maximum baseline-to-trip progression; it is not a failure probability or trained-model attribution.",
    ]
    return EquipmentInvestigation(
        asset=AssetSummary(
            asset_id="asset-he-3301",
            tag=info["Equipment Tag"],
            name=info["Equipment Name"],
            plant_id="ZCU",
            plant_name=info["Plant / Unit"],
            equipment_family="HEAT_EXCHANGER",
            equipment_type=info["Equipment Type"],
            equipment_class=info["Equipment Class"],
            discipline=info["Discipline"],
            criticality=info["Criticality"],
            monitoring_method="Weekly condition workbook and hourly DCS operating context",
        ),
        signals=signals,
        assessments=assessments,
        transitions=transitions,
        alert=AlertEvent(
            alert_id=ALERT_ID,
            asset_id="asset-he-3301",
            model_id="engineering-limits-weekly",
            policy_id=policy["policy_id"],
            first_signal_at=first.timestamp.isoformat(),
            opened_at=opened.timestamp.isoformat(),
            closed_at=None,
            status="RECOVERY_MONITORING",
            closure_reason=None,
            highest_severity="CRITICAL",
            highest_severity_rank=4,
            peak_anomaly_score=peak.score,
            peak_score_at=peak.timestamp.isoformat(),
            last_evidence_at=assessments[-1].timestamp.isoformat(),
            duration_hours=(recovered.timestamp - opened.timestamp).total_seconds() / 3600,
            primary_driver="tube_dp",
            breached_signals=peak.breached_signals,
        ),
        events=events,
        source_actions=source_actions(presentation),
        operating_states=states,
        quality_issues=issues,
        reported_downtime_hours=summary["Total Downtime (hours)"],
        reported_production_loss_tonnes=summary["Production Loss (ton)"],
    )


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--root", type=Path, default=ROOT)
    args = parser.parse_args()
    result = ingest(args.root)
    atomic_write_json(args.root / "data/normalized/he_3301/equipment.json", result)
    print(
        f"HE-3301: {len(result.assessments)} weekly assessments, 720 hourly observations, {len(result.source_actions)} source actions"
    )


if __name__ == "__main__":
    main()
