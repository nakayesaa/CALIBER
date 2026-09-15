#!/usr/bin/env python3
"""Materialize the source-anchored KO-3201 canonical dataset.

The command verifies immutable raw-file checksums before parsing, writes outputs
atomically, and produces a machine-readable validation report.  It deliberately
does not calculate anomaly events; those belong to the next analytics phase.
"""

from __future__ import annotations

import argparse
import csv
import json
import sys
from collections.abc import Sequence
from datetime import datetime, timedelta
from pathlib import Path
from typing import Any, TypeVar
from zoneinfo import ZoneInfo

import pandas as pd
import yaml
from openpyxl import load_workbook
from pptx import Presentation
from pydantic import BaseModel

REPOSITORY_ROOT = Path(__file__).resolve().parents[1]
if str(REPOSITORY_ROOT) not in sys.path:
    sys.path.insert(0, str(REPOSITORY_ROOT))

from services.api.app.schemas.canonical import (
    Action,
    ActionStatus,
    ActionType,
    Asset,
    DirectionOfConcern,
    EffectivenessCheck,
    EffectivenessResult,
    Evidence,
    EvidenceType,
    HealthState,
    Hypothesis,
    HypothesisStatus,
    Incident,
    OperatingMode,
    OperatingPeriod,
    ProductionObservation,
    QualityFlag,
    QualityIssue,
    RCACase,
    SignalDefinition,
    SignalObservation,
    SourceType,
    VerificationStatus,
    WorkflowStatus,
)
from services.api.app.services.data.normalization import (
    canonical_header,
    localize_source_timestamp,
    normalize_workflow_status,
    optional_text,
)
from services.api.app.services.data.taxonomy import (
    normalize_incident_label,
)
from services.api.app.services.file_io import (
    atomic_write_json as write_json,
)
from services.api.app.services.file_io import (
    file_sha256 as sha256_file,
)

T = TypeVar("T", bound=BaseModel)
ASSET_ID = "asset-ko-3201"
INCIDENT_ID = "incident-0002"
RCA_CASE_ID = "rca-ko-3201-2026-04-29"
SOURCE_TIMEZONE = "Asia/Jakarta"
INGESTION_VERSION = "ko_3201_ingestion_v1.0.0"
INGESTED_AT = datetime(2026, 9, 10, 0, 0, tzinfo=ZoneInfo(SOURCE_TIMEZONE))


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--root",
        type=Path,
        default=REPOSITORY_ROOT,
        help="Repository root (defaults to the script's parent repository)",
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=None,
        help="Canonical output directory (defaults to data/normalized/ko_3201)",
    )
    return parser.parse_args()


def load_and_verify_manifest(root: Path) -> tuple[dict[str, Any], dict[str, Path]]:
    manifest_path = root / "data/catalog/source_manifest.yaml"
    manifest = yaml.safe_load(manifest_path.read_text(encoding="utf-8"))
    sources: dict[str, Path] = {}
    errors: list[str] = []
    for source in manifest["sources"]:
        path = root / source["local_path"]
        if not path.is_file():
            errors.append(f"Missing source: {path}")
            continue
        actual_size = path.stat().st_size
        actual_hash = sha256_file(path)
        if actual_size != source["size_bytes"]:
            errors.append(
                f"Size mismatch for {path.name}: {actual_size} != {source['size_bytes']}"
            )
        if actual_hash != source["sha256"]:
            errors.append(f"SHA-256 mismatch for {path.name}")
        sources[source["source_key"]] = path
    if errors:
        raise RuntimeError("Raw-source verification failed:\n- " + "\n- ".join(errors))
    return manifest, sources


def scalar(value: Any) -> Any:
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return None
    return value


def model_to_row(model: BaseModel) -> dict[str, Any]:
    row = model.model_dump(mode="json")
    for key, value in row.items():
        if isinstance(value, (list, dict)):
            row[key] = json.dumps(value, separators=(",", ":"), sort_keys=True)
    return row


def write_models_csv(path: Path, models: Sequence[T]) -> None:
    if not models:
        raise ValueError(f"Refusing to write headerless empty output: {path}")
    rows = [model_to_row(model) for model in models]
    temporary = path.with_suffix(path.suffix + ".tmp")
    path.parent.mkdir(parents=True, exist_ok=True)
    with temporary.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)
    temporary.replace(path)


def equipment_metadata(path: Path) -> dict[str, str]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook["Equipment Info"]
    values: dict[str, str] = {}
    for row in sheet.iter_rows(min_row=1, max_col=2, values_only=True):
        key, value = row
        if key and value is not None:
            values[canonical_header(key)] = canonical_header(value)
    workbook.close()
    return values


def build_asset(equipment_path: Path) -> Asset:
    metadata = equipment_metadata(equipment_path)
    if metadata.get("Equipment Tag") != "KO-3201":
        raise ValueError("Equipment workbook is not the expected KO-3201 source")
    return Asset(
        asset_id=ASSET_ID,
        tag="KO-3201",
        name=metadata["Equipment Name"],
        plant_id="ZCU",
        plant_name="Zeta Cracker Unit",
        equipment_family="COMPRESSOR",
        equipment_type=metadata["Equipment Type"],
        equipment_class=metadata["Equipment Class"],
        discipline="ROTATING",
        criticality=metadata["Criticality"].upper(),
        design_life=metadata.get("Design Life"),
        monitoring_method=metadata.get("Monitoring Method"),
        active=True,
        source_reference="equipment_performance_ko_3201.xlsx#Equipment Info!A4:B12",
    )


def load_signal_mapping(root: Path) -> pd.DataFrame:
    mapping = pd.read_csv(root / "data/catalog/signal_mapping.csv")
    required = {
        "source_key",
        "source_sheet",
        "source_column",
        "signal_id",
        "canonical_name",
        "unit",
    }
    missing = required - set(mapping.columns)
    if missing:
        raise ValueError(f"Signal mapping is missing columns: {sorted(missing)}")
    if mapping["signal_id"].duplicated().any():
        duplicates = mapping.loc[mapping["signal_id"].duplicated(), "signal_id"].tolist()
        raise ValueError(f"Duplicate signal IDs: {duplicates}")
    return mapping


def build_signal_definitions(mapping: pd.DataFrame) -> list[SignalDefinition]:
    definitions: list[SignalDefinition] = []
    for record in mapping.to_dict(orient="records"):
        definitions.append(
            SignalDefinition(
                signal_id=record["signal_id"],
                asset_id=ASSET_ID,
                canonical_name=record["canonical_name"],
                source_name=record["source_column"],
                description=record["description"],
                unit=record["unit"],
                measurement_type=record["measurement_type"],
                direction_of_concern=DirectionOfConcern(record["direction_of_concern"]),
                typical_value=scalar(record["typical_value"]),
                alarm_limit=scalar(record["alarm_limit"]),
                trip_limit=scalar(record["trip_limit"]),
                source_key=record["source_key"],
                source_sheet=record["source_sheet"],
                source_cadence=record["cadence"],
                quality_status=record["quality_status"],
            )
        )
    return definitions


def quality_for_definition(definition: SignalDefinition) -> QualityFlag:
    mapping = {
        "UNIT_CONFLICT": QualityFlag.UNIT_CONFLICT,
        "AMBIGUOUS_MEASUREMENT": QualityFlag.AMBIGUOUS_MEASUREMENT,
        "SOURCE_DISAGREEMENT": QualityFlag.SOURCE_DISAGREEMENT,
        "THRESHOLD_CONTEXT_REQUIRED": QualityFlag.SOURCE_DISAGREEMENT,
    }
    return mapping.get(definition.quality_status, QualityFlag.TIMEZONE_ASSUMED)


def production_records(
    path: Path,
    definitions: Sequence[SignalDefinition],
) -> tuple[list[SignalObservation], list[ProductionObservation], list[OperatingPeriod], dict[str, Any]]:
    frame = pd.read_excel(path, sheet_name="Sheet2")
    frame.columns = [canonical_header(column) for column in frame.columns]
    frame["Timestamp"] = pd.to_datetime(frame["Timestamp"], errors="raise")
    frame = frame.sort_values("Timestamp", kind="stable").reset_index(drop=True)

    duplicate_count = int(frame["Timestamp"].duplicated().sum())
    deltas = frame["Timestamp"].diff().dropna()
    non_hourly_count = int((deltas != pd.Timedelta(hours=1)).sum())
    if duplicate_count or non_hourly_count:
        raise ValueError(
            f"Production cadence validation failed: duplicates={duplicate_count}, "
            f"non_hourly={non_hourly_count}"
        )

    by_source_column = {
        definition.source_name: definition
        for definition in definitions
        if definition.source_key == "production_ko_3201"
    }
    expected_columns = set(by_source_column) | {"Timestamp"}
    missing = expected_columns - set(frame.columns)
    if missing:
        raise ValueError(f"Production workbook missing columns: {sorted(missing)}")

    signal_observations: list[SignalObservation] = []
    production_observations: list[ProductionObservation] = []
    production_metrics = {"feed_rate", "plant_rate"}
    statuses = frame["RUN_STATUS"].astype(str).str.strip().str.upper()

    previous_status: str | None = None
    for index, row in frame.iterrows():
        timestamp = localize_source_timestamp(row["Timestamp"].to_pydatetime(), SOURCE_TIMEZONE)
        status = statuses.iloc[index]
        if status == "OFF":
            mode = OperatingMode.OFFLINE_TRIP
            health = HealthState.TRIP
        elif previous_status == "OFF" and status == "ON":
            mode = OperatingMode.RESTART
            health = HealthState.RECOVERY_MONITORING
        else:
            mode = OperatingMode.RUNNING_STEADY
            health = HealthState.UNKNOWN
        excel_row = index + 2

        for source_column, definition in by_source_column.items():
            if source_column == "RUN_STATUS":
                continue
            value = float(row[source_column])
            reference = f"production_data_ko_3201.xlsx#Sheet2!{source_column}:R{excel_row}"
            if definition.canonical_name in production_metrics:
                production_observations.append(
                    ProductionObservation(
                        observation_id=f"prod-hourly-{index + 1:04d}-{definition.canonical_name}",
                        asset_id=ASSET_ID,
                        timestamp=timestamp,
                        metric_name=definition.canonical_name,
                        value=value,
                        unit=definition.unit,
                        operating_mode=mode,
                        source_type=SourceType.OBSERVED_ANCHOR,
                        source_reference=reference,
                        quality_flag=QualityFlag.TIMEZONE_ASSUMED,
                        timezone_assumption=SOURCE_TIMEZONE,
                        ingestion_version=INGESTION_VERSION,
                        ingested_at=INGESTED_AT,
                    )
                )
            else:
                signal_observations.append(
                    SignalObservation(
                        observation_id=f"signal-hourly-{index + 1:04d}-{definition.signal_id}",
                        asset_id=ASSET_ID,
                        signal_id=definition.signal_id,
                        timestamp=timestamp,
                        value=value,
                        operating_mode=mode,
                        health_state=health,
                        source_type=SourceType.OBSERVED_ANCHOR,
                        source_reference=reference,
                        source_cadence="HOURLY",
                        quality_flag=quality_for_definition(definition),
                        timezone_assumption=SOURCE_TIMEZONE,
                        ingestion_version=INGESTION_VERSION,
                        ingested_at=INGESTED_AT,
                    )
                )
        previous_status = status

    operating_periods = collapse_operating_periods(frame)
    profile = {
        "rows": len(frame),
        "first_timestamp": str(frame["Timestamp"].iloc[0]),
        "last_timestamp": str(frame["Timestamp"].iloc[-1]),
        "duplicate_timestamps": duplicate_count,
        "non_hourly_intervals": non_hourly_count,
        "on_rows": int((statuses == "ON").sum()),
        "off_rows": int((statuses == "OFF").sum()),
    }
    return signal_observations, production_observations, operating_periods, profile


def collapse_operating_periods(frame: pd.DataFrame) -> list[OperatingPeriod]:
    timestamps = [
        localize_source_timestamp(value.to_pydatetime(), SOURCE_TIMEZONE)
        for value in frame["Timestamp"]
    ]
    statuses = frame["RUN_STATUS"].astype(str).str.strip().str.upper().tolist()
    raw_groups: list[tuple[int, int, str]] = []
    group_start = 0
    for index in range(1, len(statuses)):
        if statuses[index] != statuses[group_start]:
            raw_groups.append((group_start, index, statuses[group_start]))
            group_start = index
    raw_groups.append((group_start, len(statuses), statuses[group_start]))

    periods: list[OperatingPeriod] = []
    sequence = 1
    for group_index, (start, end, status) in enumerate(raw_groups):
        start_time = timestamps[start]
        end_time = timestamps[end] if end < len(timestamps) else timestamps[-1] + timedelta(hours=1)
        if status == "OFF":
            mode = OperatingMode.OFFLINE_TRIP
            segments = [(start_time, end_time, mode)]
        elif group_index > 0 and raw_groups[group_index - 1][2] == "OFF":
            restart_end = min(start_time + timedelta(hours=1), end_time)
            segments = [(start_time, restart_end, OperatingMode.RESTART)]
            if restart_end < end_time:
                segments.append((restart_end, end_time, OperatingMode.RUNNING_STEADY))
        else:
            segments = [(start_time, end_time, OperatingMode.RUNNING_STEADY)]

        for segment_start, segment_end, mode in segments:
            periods.append(
                OperatingPeriod(
                    operating_period_id=f"operating-period-{sequence:03d}",
                    asset_id=ASSET_ID,
                    started_at=segment_start,
                    ended_at=segment_end,
                    operating_mode=mode,
                    source_status=status,
                    source_type=SourceType.DERIVED_FEATURE,
                    source_reference="production_data_ko_3201.xlsx#Sheet2!RUN_STATUS",
                    timezone_assumption=SOURCE_TIMEZONE,
                )
            )
            sequence += 1
    return periods


def weekly_condition_records(
    path: Path,
    definitions: Sequence[SignalDefinition],
) -> tuple[list[SignalObservation], dict[str, Any], pd.DataFrame]:
    frame = pd.read_excel(path, sheet_name="Condition History")
    frame.columns = [canonical_header(column) for column in frame.columns]
    frame["Date"] = pd.to_datetime(frame["Date"], errors="raise")
    frame = frame.sort_values("Date", kind="stable").reset_index(drop=True)

    duplicate_count = int(frame["Date"].duplicated().sum())
    deltas = frame["Date"].diff().dropna()
    non_weekly_count = int((deltas != pd.Timedelta(days=7)).sum())
    if duplicate_count or non_weekly_count:
        raise ValueError(
            f"Weekly cadence validation failed: duplicates={duplicate_count}, "
            f"non_weekly={non_weekly_count}"
        )

    by_source_column = {
        definition.source_name: definition
        for definition in definitions
        if definition.source_key == "equipment_performance_ko_3201"
    }
    missing = set(by_source_column) - set(frame.columns)
    if missing:
        raise ValueError(f"Condition History missing columns: {sorted(missing)}")

    observations: list[SignalObservation] = []
    for index, row in frame.iterrows():
        timestamp = localize_source_timestamp(row["Date"].to_pydatetime(), SOURCE_TIMEZONE)
        source_health = canonical_header(row["Health Status"]).upper()
        health = HealthState(source_health)
        mode = (
            OperatingMode.OFFLINE_TRIP
            if health == HealthState.TRIP
            else OperatingMode.RUNNING_STEADY
        )
        excel_row = index + 2
        for source_column, definition in by_source_column.items():
            observations.append(
                SignalObservation(
                    observation_id=f"signal-weekly-{index + 1:03d}-{definition.signal_id}",
                    asset_id=ASSET_ID,
                    signal_id=definition.signal_id,
                    timestamp=timestamp,
                    value=float(row[source_column]),
                    operating_mode=mode,
                    health_state=health,
                    source_type=SourceType.OBSERVED_ANCHOR,
                    source_reference=(
                        "equipment_performance_ko_3201.xlsx#Condition History!"
                        f"{source_column}:R{excel_row}"
                    ),
                    source_cadence="WEEKLY",
                    quality_flag=quality_for_definition(definition),
                    timezone_assumption=SOURCE_TIMEZONE,
                    ingestion_version=INGESTION_VERSION,
                    ingested_at=INGESTED_AT,
                )
            )
    profile = {
        "rows": len(frame),
        "first_date": str(frame["Date"].iloc[0].date()),
        "last_date": str(frame["Date"].iloc[-1].date()),
        "duplicate_dates": duplicate_count,
        "non_weekly_intervals": non_weekly_count,
        "health_status_counts": frame["Health Status"].value_counts().to_dict(),
    }
    return observations, profile, frame


def incident_records(
    path: Path,
    taxonomy: dict[str, Any],
) -> tuple[list[Incident], list[BaseModel], dict[str, Any]]:
    frame = pd.read_excel(path, sheet_name="Incident Database", header=2)
    frame.columns = [canonical_header(column) for column in frame.columns]
    required = {
        "Serial No",
        "MTO No.",
        "AR No.",
        "Plant",
        "Tag Number",
        "Eq. Class",
        "Date of Occur.",
        "Risk Case Title",
        "Highest Impact",
        "Pre-Risk",
        "Risk Score",
        "PIC (RCA)",
        "Overall Status",
        "Discipline",
        "Eq. Type",
        "Component",
        "F Mechanism",
        "Downtime (hrs)",
        "Act. Loss (k US$)",
        "Pot. Loss (k US$)",
        "Total Loss (k US$)",
        "RCA Due Date",
    }
    missing = required - set(frame.columns)
    if missing:
        raise ValueError(f"Incident Database missing columns: {sorted(missing)}")

    incidents: list[Incident] = []
    labels: list[BaseModel] = []
    for index, row in frame.iterrows():
        serial = int(row["Serial No"])
        asset_tag = canonical_header(row["Tag Number"])
        occurred_at = localize_source_timestamp(row["Date of Occur."], SOURCE_TIMEZONE)
        rca_due = optional_text(row["RCA Due Date"])
        if rca_due is not None:
            rca_due = str(pd.Timestamp(rca_due).date())
        incident = Incident(
            incident_id=f"incident-{serial:04d}",
            source_serial=serial,
            mto_number=optional_text(row["MTO No."]),
            ar_number=optional_text(row["AR No."]),
            asset_tag=asset_tag,
            asset_id=ASSET_ID if asset_tag == "KO-3201" else None,
            is_anchor_asset=asset_tag == "KO-3201",
            plant_id=canonical_header(row["Plant"]),
            occurred_at=occurred_at,
            title=canonical_header(row["Risk Case Title"]),
            highest_impact=canonical_header(row["Highest Impact"]),
            equipment_class=canonical_header(row["Eq. Class"]),
            pre_risk=canonical_header(row["Pre-Risk"]),
            risk_score=float(row["Risk Score"]),
            owner=canonical_header(row["PIC (RCA)"]),
            workflow_status=normalize_workflow_status(row["Overall Status"]),
            discipline_raw=canonical_header(row["Discipline"]),
            equipment_type_raw=canonical_header(row["Eq. Type"]),
            component_raw=canonical_header(row["Component"]),
            mechanism_raw=canonical_header(row["F Mechanism"]),
            downtime_hours=float(row["Downtime (hrs)"]),
            actual_loss_kusd=float(row["Act. Loss (k US$)"]),
            potential_loss_kusd=float(row["Pot. Loss (k US$)"]),
            total_loss_kusd=float(row["Total Loss (k US$)"]),
            rca_due_date=rca_due,
            source_reference=f"incident_database.xlsx#Incident Database!R{index + 4}",
        )
        incidents.append(incident)
        labels.append(normalize_incident_label(incident, taxonomy))

    anchor_matches = [incident for incident in incidents if incident.asset_tag == "KO-3201"]
    if len(anchor_matches) != 1:
        raise ValueError(f"Expected one KO-3201 incident, found {len(anchor_matches)}")
    if anchor_matches[0].incident_id != INCIDENT_ID:
        raise ValueError(
            f"KO-3201 source serial changed: {anchor_matches[0].incident_id} != {INCIDENT_ID}"
        )
    profile = {
        "rows": len(frame),
        "ko_3201_matches": len(anchor_matches),
        "downtime_hours_total": round(sum(item.downtime_hours for item in incidents), 3),
        "total_loss_kusd": round(sum(item.total_loss_kusd for item in incidents), 3),
    }
    return incidents, labels, profile


def presentation_text(path: Path) -> list[str]:
    presentation = Presentation(path)
    slides: list[str] = []
    for slide in presentation.slides:
        slides.append(
            "\n".join(
                shape.text.strip()
                for shape in slide.shapes
                if hasattr(shape, "text") and shape.text.strip()
            )
        )
    return slides


def validate_rca_source(path: Path) -> list[str]:
    slides = presentation_text(path)
    required_markers = {
        1: "AR-2026-ZCU-0142",
        3: "29-Apr-2026 06:40",
        7: "ROOT CAUSE:",
        9: "Install online water-in-oil sensor",
        11: "PRODUCTION LOSS",
    }
    if len(slides) != 11:
        raise ValueError(f"Expected 11 KO-3201 RCA slides, found {len(slides)}")
    for slide_number, marker in required_markers.items():
        if marker not in slides[slide_number - 1]:
            raise ValueError(f"RCA slide {slide_number} missing marker {marker!r}")
    return slides


def rca_case() -> RCACase:
    return RCACase(
        rca_case_id=RCA_CASE_ID,
        incident_id=INCIDENT_ID,
        asset_id=ASSET_ID,
        problem_statement=(
            "Cracked Gas Compressor KO-3201 tripped on high-high radial vibration "
            "on 29-Apr-2026 due to source-reported journal-bearing babbitt distress "
            "driven by water-contaminated lube oil."
        ),
        source_reported_root_cause=(
            "Journal-bearing babbitt distress caused by degraded oil-film strength "
            "from water ingress via a leaking lube-oil cooler tube, undetected due "
            "to absence of online water-in-oil monitoring."
        ),
        workflow_status=WorkflowStatus.CA_PA_EXECUTION,
        owner="REL-05",
        occurred_at=localize_source_timestamp("2026-04-29 06:40:00", SOURCE_TIMEZONE),
        reported_at=localize_source_timestamp("2026-04-30 00:00:00", SOURCE_TIMEZONE),
        source_type=SourceType.OBSERVED_ANCHOR,
        source_reference="rca_ko_3201_high_radial_vibration_trip.pptx#slides-1-11",
        app_confirmation_status="SOURCE_REPORTED_NOT_APP_CONFIRMED",
    )


def rca_evidence() -> list[Evidence]:
    tz = SOURCE_TIMEZONE
    return [
        Evidence(
            evidence_id="evidence-ko-001",
            rca_case_id=RCA_CASE_ID,
            incident_id=INCIDENT_ID,
            asset_id=ASSET_ID,
            evidence_type=EvidenceType.SENSOR,
            title="Weekly four-signal degradation history",
            observed_at=localize_source_timestamp("2026-04-29", tz),
            content_summary=(
                "Condition workbook shows vibration and water rising, oil pressure "
                "falling, and bearing temperature rising from normal through trip."
            ),
            verification_status=VerificationStatus.AVAILABLE_UNVERIFIED,
            evidence_grade="A_SOURCE_RECORD",
            source_type=SourceType.OBSERVED_ANCHOR,
            source_reference="equipment_performance_ko_3201.xlsx#Condition History!R2:R27",
            quality_flag=QualityFlag.TIMEZONE_ASSUMED,
        ),
        Evidence(
            evidence_id="evidence-ko-002",
            rca_case_id=RCA_CASE_ID,
            incident_id=INCIDENT_ID,
            asset_id=ASSET_ID,
            evidence_type=EvidenceType.SENSOR,
            title="High-high vibration trip reference",
            observed_at=localize_source_timestamp("2026-04-29 06:40:00", tz),
            content_summary="RCA chronology reports DE vibration reached 75 micron and tripped the compressor.",
            verification_status=VerificationStatus.REFERENCED_NOT_AVAILABLE,
            evidence_grade="B_SOURCE_REPORT",
            source_type=SourceType.OBSERVED_ANCHOR,
            source_reference="rca_ko_3201_high_radial_vibration_trip.pptx#slide-3",
            quality_flag=QualityFlag.MISSING_EVIDENCE,
        ),
        Evidence(
            evidence_id="evidence-ko-003",
            rca_case_id=RCA_CASE_ID,
            incident_id=INCIDENT_ID,
            asset_id=ASSET_ID,
            evidence_type=EvidenceType.LAB,
            title="Lube-oil sample water-content reference",
            observed_at=localize_source_timestamp("2026-04-29 08:10:00", tz),
            content_summary="RCA chronology reports 1,800 ppm water versus specification below 500 ppm.",
            verification_status=VerificationStatus.REFERENCED_NOT_AVAILABLE,
            evidence_grade="B_SOURCE_REPORT",
            source_type=SourceType.OBSERVED_ANCHOR,
            source_reference="rca_ko_3201_high_radial_vibration_trip.pptx#slides-3-and-6",
            quality_flag=QualityFlag.MISSING_EVIDENCE,
        ),
        Evidence(
            evidence_id="evidence-ko-004",
            rca_case_id=RCA_CASE_ID,
            incident_id=INCIDENT_ID,
            asset_id=ASSET_ID,
            evidence_type=EvidenceType.INSPECTION,
            title="Journal-bearing babbitt distress reference",
            observed_at=localize_source_timestamp("2026-04-29 18:00:00", tz),
            content_summary="RCA chronology reports the DE journal bearing was removed and babbitt was wiped/distressed.",
            verification_status=VerificationStatus.REFERENCED_NOT_AVAILABLE,
            evidence_grade="B_SOURCE_REPORT",
            source_type=SourceType.OBSERVED_ANCHOR,
            source_reference="rca_ko_3201_high_radial_vibration_trip.pptx#slide-3",
            quality_flag=QualityFlag.MISSING_EVIDENCE,
        ),
        Evidence(
            evidence_id="evidence-ko-005",
            rca_case_id=RCA_CASE_ID,
            incident_id=INCIDENT_ID,
            asset_id=ASSET_ID,
            evidence_type=EvidenceType.INSPECTION,
            title="Oil-cooler tube leak reference",
            observed_at=localize_source_timestamp("2026-04-29 18:00:00", tz),
            content_summary="RCA chronology states that a lube-oil cooler tube leak was confirmed.",
            verification_status=VerificationStatus.REFERENCED_NOT_AVAILABLE,
            evidence_grade="B_SOURCE_REPORT",
            source_type=SourceType.OBSERVED_ANCHOR,
            source_reference="rca_ko_3201_high_radial_vibration_trip.pptx#slides-3-and-7",
            quality_flag=QualityFlag.MISSING_EVIDENCE,
        ),
        Evidence(
            evidence_id="evidence-ko-006",
            rca_case_id=RCA_CASE_ID,
            incident_id=INCIDENT_ID,
            asset_id=ASSET_ID,
            evidence_type=EvidenceType.SENSOR,
            title="No process surge reported",
            observed_at=localize_source_timestamp("2026-04-29 06:40:00", tz),
            content_summary="RCA parameter verification reports anti-surge margin maintained with no surge event.",
            verification_status=VerificationStatus.REFERENCED_NOT_AVAILABLE,
            evidence_grade="B_SOURCE_REPORT",
            source_type=SourceType.OBSERVED_ANCHOR,
            source_reference="rca_ko_3201_high_radial_vibration_trip.pptx#slide-6",
            quality_flag=QualityFlag.MISSING_EVIDENCE,
        ),
        Evidence(
            evidence_id="evidence-ko-007",
            rca_case_id=RCA_CASE_ID,
            incident_id=INCIDENT_ID,
            asset_id=ASSET_ID,
            evidence_type=EvidenceType.SENSOR,
            title="Conflicting lube-oil pressure statements",
            observed_at=localize_source_timestamp("2026-04-29", tz),
            content_summary=(
                "Weekly condition row reports 1.078 barg while RCA parameter "
                "verification reports 1.8 barg within normal band."
            ),
            verification_status=VerificationStatus.AVAILABLE_UNVERIFIED,
            evidence_grade="A_AND_B_CONFLICT",
            source_type=SourceType.OBSERVED_ANCHOR,
            source_reference=(
                "equipment_performance_ko_3201.xlsx#Condition History!R22;"
                "rca_ko_3201_high_radial_vibration_trip.pptx#slide-6"
            ),
            quality_flag=QualityFlag.SOURCE_DISAGREEMENT,
        ),
    ]


def rca_hypotheses() -> list[Hypothesis]:
    return [
        Hypothesis(
            hypothesis_id="hypothesis-ko-001",
            rca_case_id=RCA_CASE_ID,
            incident_id=INCIDENT_ID,
            asset_id=ASSET_ID,
            statement=(
                "Water contamination degraded oil-film strength, causing journal-bearing "
                "distress and high radial vibration."
            ),
            rank=1,
            confidence=0.88,
            status=HypothesisStatus.PROPOSED,
            supporting_evidence_ids=["evidence-ko-001", "evidence-ko-003", "evidence-ko-004"],
            contradicting_evidence_ids=["evidence-ko-007"],
            missing_evidence=["Signed lab report", "Bearing inspection report/photos"],
            generated_by="SOURCE_RCA_DECK",
            source_type=SourceType.OBSERVED_ANCHOR,
        ),
        Hypothesis(
            hypothesis_id="hypothesis-ko-002",
            rca_case_id=RCA_CASE_ID,
            incident_id=INCIDENT_ID,
            asset_id=ASSET_ID,
            statement="Loss of oil-cooler integrity allowed water ingress into the lube-oil system.",
            rank=2,
            confidence=0.82,
            status=HypothesisStatus.PROPOSED,
            supporting_evidence_ids=["evidence-ko-003", "evidence-ko-005"],
            contradicting_evidence_ids=[],
            missing_evidence=["Cooler leak-test report", "Repair work order"],
            generated_by="SOURCE_RCA_DECK",
            source_type=SourceType.OBSERVED_ANCHOR,
        ),
        Hypothesis(
            hypothesis_id="hypothesis-ko-003",
            rca_case_id=RCA_CASE_ID,
            incident_id=INCIDENT_ID,
            asset_id=ASSET_ID,
            statement=(
                "A different rotor or bearing mechanical fault produced vibration, "
                "with water contamination acting only as a contributing condition."
            ),
            rank=3,
            confidence=0.28,
            status=HypothesisStatus.PROPOSED,
            supporting_evidence_ids=["evidence-ko-002"],
            contradicting_evidence_ids=["evidence-ko-003", "evidence-ko-005"],
            missing_evidence=["Rotor inspection", "Bearing failure analysis"],
            generated_by="CALIBER_DETERMINISTIC_ALTERNATIVE_V1",
            source_type=SourceType.DERIVED_FEATURE,
        ),
    ]


def rca_actions() -> list[Action]:
    action_specs = [
        ("001", ActionType.CONTAINMENT, "Controlled shutdown, replace DE journal bearing, and flush lube-oil system", "REL-05", "2026-04-30", ActionStatus.COMPLETED, "Authorized shutdown record and maintenance work order", "slide-2"),
        ("002", ActionType.CORRECTIVE, "Repair leaking lube-oil cooler tube (plug and re-test)", "STA-02", "2026-04-30", ActionStatus.COMPLETED, "Cooler leak-test and repair work order", "slide-9"),
        ("003", ActionType.CORRECTIVE, "Install online water-in-oil sensor on KO-3201 lube system", "REL-05", "2026-06-30", ActionStatus.IN_PROGRESS, "Installation, calibration, and commissioning record", "slide-9"),
        ("004", ActionType.PROACTIVE, "Retube lube-oil cooler at next turnaround", "STA-02", "2026-10-31", ActionStatus.OPEN, "Turnaround work order and pressure-test result", "slide-9"),
        ("005", ActionType.PROACTIVE, "Apply online water-in-oil monitoring to KO-3202 and KO-3203", "REL-05", "2026-09-30", ActionStatus.OPEN, "Fleet rollout and commissioning records", "slide-9"),
        ("006", ActionType.PREVENTIVE, "Tighten KO-3201 vibration alert to 45 micron and shorten trend-review cadence", "REL-05", "2026-05-30", ActionStatus.OPEN, "Approved setpoint change and startup suppression test", "slide-10"),
        ("007", ActionType.PREVENTIVE, "Add lube-oil cooler differential-pressure and conductivity trends to reliability dashboard", "REL-02", "2026-06-30", ActionStatus.OPEN, "Dashboard configuration and signal validation", "slide-10"),
        ("008", ActionType.PREVENTIVE, "Establish weekly lube-oil water-content and particle-count PM", "REL-05", None, ActionStatus.OPEN, "Approved PM schedule and first completed sample", "slide-10"),
        ("009", ActionType.PREVENTIVE, "Establish six-monthly lube-oil cooler leak test", "STA-02", None, ActionStatus.OPEN, "Approved PM schedule and leak-test report", "slide-10"),
        ("010", ActionType.PREVENTIVE, "Establish daily continuous vibration trend review for KO-3201", "REL-05", None, ActionStatus.OPEN, "Approved monitoring task and review record", "slide-10"),
    ]
    actions: list[Action] = []
    for number, action_type, description, owner, due_date, status, evidence, slide in action_specs:
        completed_at = None
        if number in {"001", "002"}:
            completed_at = localize_source_timestamp("2026-04-30 14:40:00", SOURCE_TIMEZONE)
        actions.append(
            Action(
                action_id=f"action-ko-{number}",
                rca_case_id=RCA_CASE_ID,
                incident_id=INCIDENT_ID,
                asset_id=ASSET_ID,
                action_type=action_type,
                description=description,
                owner=owner,
                due_date=due_date,
                priority="P1" if number in {"001", "002", "003"} else "P2",
                status=status,
                required_evidence=evidence,
                completion_evidence_id=None,
                approved_by=None,
                completed_at=completed_at,
                source_type=SourceType.OBSERVED_ANCHOR,
                source_reference=f"rca_ko_3201_high_radial_vibration_trip.pptx#{slide}",
            )
        )
    return actions


def effectiveness_check(weekly_frame: pd.DataFrame) -> EffectivenessCheck:
    trip = weekly_frame.loc[weekly_frame["Health Status"] == "TRIP"].iloc[0]
    post = weekly_frame.loc[weekly_frame["Date"] >= pd.Timestamp("2026-05-06")]
    first_post = post.iloc[0]
    metrics = {
        "radial_vibration_micron": {"before": float(trip["DE Radial Vibration (micron)"]), "after": float(first_post["DE Radial Vibration (micron)"])},
        "water_in_oil_ppm": {"before": float(trip["Lube Oil Water Content (ppm)"]), "after": float(first_post["Lube Oil Water Content (ppm)"])},
        "lube_oil_pressure_barg": {"before": float(trip["Lube Oil Supply Press (barg)"]), "after": float(first_post["Lube Oil Supply Press (barg)"])},
        "bearing_metal_temperature_degc": {"before": float(trip["Bearing Metal Temp (°C)"]), "after": float(first_post["Bearing Metal Temp (°C)"])},
        "post_repair_normal_weeks": int((post["Health Status"] == "NORMAL").sum()),
    }
    return EffectivenessCheck(
        effectiveness_check_id="effectiveness-ko-001",
        rca_case_id=RCA_CASE_ID,
        incident_id=INCIDENT_ID,
        asset_id=ASSET_ID,
        monitoring_start=localize_source_timestamp("2026-05-06", SOURCE_TIMEZONE),
        monitoring_end=localize_source_timestamp("2026-06-03 23:59:59", SOURCE_TIMEZONE),
        baseline_window="2025-12-10 through 2026-02-04 weekly NORMAL observations",
        comparison_metrics=metrics,
        recurrence_detected=False,
        result=EffectivenessResult.INITIAL_EFFECTIVE,
        explanation=(
            "Five supplied post-repair weekly rows are NORMAL and the four anchor "
            "signals improve. Final effectiveness and risk closure still require "
            "an authorized human decision."
        ),
        approved_by=None,
        approved_at=None,
        source_type=SourceType.DERIVED_FEATURE,
        source_reference="equipment_performance_ko_3201.xlsx#Condition History!R22:R27",
    )


def quality_issues() -> list[QualityIssue]:
    return [
        QualityIssue(
            issue_id="quality-ko-001",
            severity="HIGH",
            entity_type="SIGNAL_DEFINITION",
            entity_id="sig-ko-vibration-hourly",
            quality_flag=QualityFlag.UNIT_CONFLICT,
            description="Hourly PI metadata reports vibration in mm/s; weekly/RCA records use micron. No conversion is applied.",
            source_references=["production_data_ko_3201.xlsx#PI Tag!R4", "equipment_performance_ko_3201.xlsx#Equipment Info!D5"],
            resolution_status="OPEN",
            blocks_scoring=True,
        ),
        QualityIssue(
            issue_id="quality-ko-002",
            severity="MEDIUM",
            entity_type="SIGNAL_OBSERVATION",
            entity_id="radial_vibration_de",
            quality_flag=QualityFlag.CADENCE_DIFFERENCE,
            description="Hourly peak is 72.67 while weekly/RCA context reports approximately 75 to 76.5; retain separate cadence/source representations.",
            source_references=["production_data_ko_3201.xlsx#Sheet2!KO3201_VIB", "equipment_performance_ko_3201.xlsx#Condition History!R22", "rca_ko_3201_high_radial_vibration_trip.pptx#slide-3"],
            resolution_status="ACCEPTED_SOURCE_CONTEXT",
            blocks_scoring=False,
        ),
        QualityIssue(
            issue_id="quality-ko-003",
            severity="MEDIUM",
            entity_type="EVIDENCE",
            entity_id="water_in_oil",
            quality_flag=QualityFlag.SOURCE_DISAGREEMENT,
            description="Weekly trip row reports 1,530 ppm; RCA chronology reports a lab sample of 1,800 ppm.",
            source_references=["equipment_performance_ko_3201.xlsx#Condition History!R22", "rca_ko_3201_high_radial_vibration_trip.pptx#slide-3"],
            resolution_status="OPEN_KEEP_BOTH",
            blocks_scoring=False,
        ),
        QualityIssue(
            issue_id="quality-ko-004",
            severity="HIGH",
            entity_type="EVIDENCE",
            entity_id="lube_oil_pressure",
            quality_flag=QualityFlag.SOURCE_DISAGREEMENT,
            description="Weekly trip row reports 1.078 barg, but RCA parameter verification reports 1.8 barg and marks pressure good.",
            source_references=["equipment_performance_ko_3201.xlsx#Condition History!R22", "rca_ko_3201_high_radial_vibration_trip.pptx#slide-6"],
            resolution_status="OPEN_KEEP_BOTH",
            blocks_scoring=False,
        ),
        QualityIssue(
            issue_id="quality-ko-005",
            severity="HIGH",
            entity_type="SIGNAL_DEFINITION",
            entity_id="sig-ko-vibration-weekly",
            quality_flag=QualityFlag.SOURCE_DISAGREEMENT,
            description="Equipment sheet lists a 45-micron alarm; RCA states the historical alert was 60 and recommends tightening it to 45. Effective-date context is required.",
            source_references=["equipment_performance_ko_3201.xlsx#Equipment Info!D5", "rca_ko_3201_high_radial_vibration_trip.pptx#slides-7-and-10"],
            resolution_status="OPEN_THRESHOLD_VERSION_REQUIRED",
            blocks_scoring=False,
        ),
        QualityIssue(
            issue_id="quality-ko-006",
            severity="MEDIUM",
            entity_type="SIGNAL_DEFINITION",
            entity_id="sig-ko-bearing-process-temp-hourly",
            quality_flag=QualityFlag.AMBIGUOUS_MEASUREMENT,
            description="Hourly PI description combines bearing/process temperature and cannot be assumed identical to weekly bearing-metal temperature.",
            source_references=["production_data_ko_3201.xlsx#PI Tag!R5", "equipment_performance_ko_3201.xlsx#Condition History!F:F"],
            resolution_status="OPEN",
            blocks_scoring=True,
        ),
        QualityIssue(
            issue_id="quality-ko-007",
            severity="MEDIUM",
            entity_type="DATASET",
            entity_id="ko_3201_timestamps",
            quality_flag=QualityFlag.TIMEZONE_ASSUMED,
            description="Source timestamps do not declare a timezone; Asia/Jakarta is attached as an explicit prototype assumption.",
            source_references=["production_data_ko_3201.xlsx", "equipment_performance_ko_3201.xlsx", "incident_database.xlsx", "rca_ko_3201_high_radial_vibration_trip.pptx"],
            resolution_status="DOCUMENTED_ASSUMPTION",
            blocks_scoring=False,
        ),
        QualityIssue(
            issue_id="quality-ko-008",
            severity="HIGH",
            entity_type="RCA_CASE",
            entity_id=RCA_CASE_ID,
            quality_flag=QualityFlag.MISSING_EVIDENCE,
            description="Primary lab, bearing inspection, cooler leak-test, machinery log, and work-order artifacts referenced by the RCA deck are not supplied.",
            source_references=["rca_ko_3201_high_radial_vibration_trip.pptx#slides-3-10"],
            resolution_status="OPEN_HUMAN_EVIDENCE_REQUIRED",
            blocks_scoring=False,
        ),
        QualityIssue(
            issue_id="quality-ko-009",
            severity="LOW",
            entity_type="ASSET",
            entity_id=ASSET_ID,
            quality_flag=QualityFlag.SOURCE_DISAGREEMENT,
            description="Equipment information lists five-year bearing design life; RCA parameter verification states six years.",
            source_references=["equipment_performance_ko_3201.xlsx#Equipment Info!B11", "rca_ko_3201_high_radial_vibration_trip.pptx#slide-6"],
            resolution_status="OPEN_EXCLUDED_FROM_SCORING",
            blocks_scoring=False,
        ),
    ]


def check(name: str, actual: Any, expected: Any) -> dict[str, Any]:
    return {
        "name": name,
        "status": "PASS" if actual == expected else "FAIL",
        "actual": actual,
        "expected": expected,
    }


def run(root: Path, output: Path) -> dict[str, Any]:
    root = root.resolve()
    output = output.resolve()
    manifest, sources = load_and_verify_manifest(root)
    mapping = load_signal_mapping(root)
    taxonomy = yaml.safe_load(
        (root / "data/catalog/incident_taxonomy.yaml").read_text(encoding="utf-8")
    )

    asset = build_asset(sources["equipment_performance_ko_3201"])
    definitions = build_signal_definitions(mapping)
    hourly_signals, production, periods, production_profile = production_records(
        sources["production_ko_3201"], definitions
    )
    weekly_signals, weekly_profile, weekly_frame = weekly_condition_records(
        sources["equipment_performance_ko_3201"], definitions
    )
    incidents, labels, incident_profile = incident_records(
        sources["incident_database"], taxonomy
    )
    slides = validate_rca_source(sources["rca_ko_3201"])
    case = rca_case()
    evidence = rca_evidence()
    hypotheses = rca_hypotheses()
    actions = rca_actions()
    effectiveness = effectiveness_check(weekly_frame)
    issues = quality_issues()
    signals = hourly_signals + weekly_signals

    write_models_csv(output / "assets.csv", [asset])
    write_models_csv(output / "signal_definitions.csv", definitions)
    write_models_csv(output / "signal_observations.csv", signals)
    write_models_csv(output / "production_observations.csv", production)
    write_models_csv(output / "operating_periods.csv", periods)
    write_models_csv(output / "incidents.csv", incidents)
    write_models_csv(output / "incident_labels.csv", labels)
    write_models_csv(output / "rca_cases.csv", [case])
    write_models_csv(output / "evidence.csv", evidence)
    write_models_csv(output / "hypotheses.csv", hypotheses)
    write_models_csv(output / "actions.csv", actions)
    write_models_csv(output / "effectiveness_checks.csv", [effectiveness])
    write_models_csv(output / "quality_issues.csv", issues)

    checks = [
        check("verified_raw_sources", len(sources), 5),
        check("production_hourly_rows", production_profile["rows"], 720),
        check("production_duplicate_timestamps", production_profile["duplicate_timestamps"], 0),
        check("production_non_hourly_intervals", production_profile["non_hourly_intervals"], 0),
        check("production_offline_rows", production_profile["off_rows"], 32),
        check("weekly_condition_rows", weekly_profile["rows"], 26),
        check("weekly_duplicate_dates", weekly_profile["duplicate_dates"], 0),
        check("weekly_non_weekly_intervals", weekly_profile["non_weekly_intervals"], 0),
        check("historical_incident_rows", incident_profile["rows"], 380),
        check("ko_3201_incident_matches", incident_profile["ko_3201_matches"], 1),
        check("rca_slide_count", len(slides), 11),
        check("signal_definition_count", len(definitions), 11),
        check("signal_observation_count", len(signals), 2984),
        check("production_observation_count", len(production), 1440),
        check("operating_period_count", len(periods), 4),
        check("evidence_count", len(evidence), 7),
        check("action_count", len(actions), 10),
    ]
    failed = [item for item in checks if item["status"] == "FAIL"]
    report = {
        "dataset": "KO-3201 canonical source dataset",
        "dataset_version": "1.0.0",
        "ingestion_version": INGESTION_VERSION,
        "status": "FAIL" if failed else "PASS_WITH_DOCUMENTED_WARNINGS",
        "generated_at": INGESTED_AT.isoformat(),
        "source_timezone_assumption": SOURCE_TIMEZONE,
        "source_manifest_version": manifest["manifest_version"],
        "profiles": {
            "production": production_profile,
            "weekly_condition": weekly_profile,
            "incidents": incident_profile,
        },
        "output_counts": {
            "assets": 1,
            "signal_definitions": len(definitions),
            "signal_observations": len(signals),
            "production_observations": len(production),
            "operating_periods": len(periods),
            "incidents": len(incidents),
            "incident_labels": len(labels),
            "rca_cases": 1,
            "evidence": len(evidence),
            "hypotheses": len(hypotheses),
            "actions": len(actions),
            "effectiveness_checks": 1,
            "quality_issues": len(issues),
        },
        "checks": checks,
        "warning_summary": {
            "count": len(issues),
            "blocking_for_specific_analytics": sum(issue.blocks_scoring for issue in issues),
            "note": (
                "Warnings are governed outputs. They do not invalidate source ingestion, "
                "but some representations are excluded from future anomaly scoring until resolved."
            ),
        },
    }
    write_json(output / "validation_report.json", report)
    write_json(
        output / "dataset_metadata.json",
        {
            "dataset": report["dataset"],
            "dataset_version": report["dataset_version"],
            "ingestion_version": INGESTION_VERSION,
            "asset_id": ASSET_ID,
            "asset_tag": "KO-3201",
            "primary_incident_id": INCIDENT_ID,
            "rca_case_id": RCA_CASE_ID,
            "source_manifest": "data/catalog/source_manifest.yaml",
            "lineage_policy": [item.value for item in SourceType],
            "quality_report": "validation_report.json",
        },
    )
    if failed:
        raise RuntimeError(f"Canonical validation failed: {failed}")
    return report


def main() -> None:
    args = parse_args()
    root = args.root.resolve()
    output = (
        args.output.resolve()
        if args.output is not None
        else root / "data/normalized/ko_3201"
    )
    report = run(root, output)
    print(json.dumps(report, indent=2, sort_keys=True))


if __name__ == "__main__":
    main()
